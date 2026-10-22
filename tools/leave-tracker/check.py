import openpyxl,sys
wb=openpyxl.load_workbook(sys.argv[1],data_only=True)
for name in ['Summary 2026','Summary 2027']:
    s=wb[name]; print('==',name,[s.cell(3,c).value for c in range(2,13)])
    for r in range(6,12):
        row=[s.cell(r,c).value for c in range(2,20)]
        if row[0]: print(row)
s=wb['Staff']
for r in range(5,9): print([s.cell(r,c).value for c in range(1,9)])
l=wb['Leave Log']
for r in range(5,12):
    row=[l.cell(r,c).value for c in range(1,15)]
    if row[0]: print(row)
c=wb['Calendar']
print('cal',c['B2'].value,c['E2'].value)
for r in range(4,11):
    print([c.cell(r,k).value for k in range(1,36)])
