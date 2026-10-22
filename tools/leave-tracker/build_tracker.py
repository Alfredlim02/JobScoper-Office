"""Build the JobScoper Leave Tracker 2026-2027 workbook (v2: JobScoper branding, auto Birthday Leave,
separate holiday tabs). Output is imported into Google Sheets (File > Import > Replace spreadsheet)."""
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else "JobScoper_Leave_Tracker_2026-2027.xlsx"

# ---------------------------------------------------------------- JobScoper brand theme (from jobscoper.com.my)
BLACK = "000000"
YELLOW = "FFD400"        # brand yellow
YELLOW_SOFT = "FFF8D6"   # cells Alfred types into
LILAC = "E3DDF0"
SKY = "D6EEF7"
AUTO_FILL = "F4F4F4"     # formula cells
GREY_TXT = "6C757D"
BORDER = "D9D9D9"
LEAVE = {  # fill, font
    "AL": ("FFE45C", "000000"),
    "MC": (SKY, "0B4F6C"),
    "BL": (LILAC, "4B3B7A"),
    "PH": ("111111", YELLOW),
}
WEEKEND = "E9E9E9"
STATUS = {
    "Upcoming": (SKY, "0B4F6C"),
    "New Joiner": (LILAC, "4B3B7A"),
    "Active": ("DFF3E4", "1E6B34"),
    "Resigning": ("FFE8B3", "8A5A00"),
    "Left": ("E9ECEF", "6C757D"),
}
FONT = "Nunito"          # body (rounded, like the website)
TITLE_FONT = "Baloo 2"   # headings

thin = Side(style="thin", color=BORDER)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def F(size=10, bold=False, color="212529", italic=False, name=FONT):
    return Font(name=name, size=size, bold=bold, color=color, italic=italic)


def header(ws, row, labels, col=1, fill_c=BLACK, font_c="FFFFFF", height=34):
    for i, lab in enumerate(labels):
        c = ws.cell(row, col + i, lab)
        c.font = F(10, True, font_c)
        c.fill = fill(fill_c)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = height


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(name=TITLE_FONT, bold=True, size=18, color=BLACK)
    ws.row_dimensions[1].height = 32
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F(9, italic=True, color=GREY_TXT)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def cf(ws, rng, formula, fill_c=None, font_c=None, bold=False):
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[formula], fill=fill(fill_c) if fill_c else None,
        font=Font(color=font_c, bold=bold) if font_c else None))


def status_cf(ws, rng, first_cell):
    for s, (f, t) in STATUS.items():
        cf(ws, rng, f'{first_cell}="{s}"', f, t, True)


wb = Workbook()

# ---------------------------------------------------------------- layout constants + named ranges
STAFF_FIRST, STAFF_LAST = 5, 34          # 30 staff slots
LOG_FIRST, LOG_LAST = 5, 504             # 500 leave rows
HOL_FIRST, HOL_LAST = 5, 44              # 40 holiday rows per year
YEARS = (2026, 2027)
NAMES = {}


def name(n, ref):
    NAMES[n] = ref
    return n


def staff(col):
    return name(f"Staff_{col}", f"Staff!${col}${STAFF_FIRST}:${col}${STAFF_LAST}")


def log(col):
    return name(f"LeaveLog_{col}", f"'Leave Log'!${col}${LOG_FIRST}:${col}${LOG_LAST}")


HOLS = {y: name(f"Hols_{y}", f"'Holidays {y}'!$A${HOL_FIRST}:$A${HOL_LAST}") for y in YEARS}
HOL_NAMES = {y: name(f"HolNames_{y}", f"'Holidays {y}'!$B${HOL_FIRST}:$B${HOL_LAST}") for y in YEARS}


def hols_for(date_expr):
    """Holiday list matching the year of date_expr (2027 list for 2027, otherwise 2026)."""
    return f"IF(YEAR({date_expr})=2027,{HOLS[2027]},{HOLS[2026]})"


def is_ph(date_expr):
    return f"(COUNTIF({HOLS[2026]},{date_expr})+COUNTIF({HOLS[2027]},{date_expr})>0)"


# Staff columns
S_NAME, S_POS, S_START, S_LWD, S_BDAY, S_BLOVR, S_CF, S_STATUS, S_SVC, S_NOTES = "ABCDEFGHIJ"

# ================================================================ How to use
ws = wb.active
ws.title = "How to use"
title(ws, "JobScoper Leave Tracker 2026–2027", "Owner: Alfred (People Solution Lead) · Only Alfred edits this file")
sections = [
    ("HOW TO APPLY LEAVE (Leave Log tab)", [
        ("Step 1", "Open the Leave Log tab and go to the first empty row."),
        ("Step 2", "Name → pick from dropdown.   Type → AL or MC.   From → first day.   To → last day (leave blank for 1 day)."),
        ("Step 3", "Day → Full, or AM / PM for half day.   Status → Approved.   Remarks → optional."),
        ("Step 4", "Done. Working Days and ⚠ Check fill in by themselves; Summary and Calendar update instantly."),
        ("Cancel?", "Change Status to Cancelled — don't delete the row."),
        ("Example 1", "Shamirah · AL · 2 Dec 2026 · 4 Dec 2026 · Full · Approved  →  3 working days"),
        ("Example 2", "Zack · MC · 18 Nov 2026 · (blank) · Full · Approved  →  1 day"),
        ("Example 3", "Alfred · AL · 9 Dec 2026 · (blank) · PM · Approved  →  0.5 day"),
    ]),
    ("PEOPLE (Staff tab)", [
        ("New joiner", "Add Name, Position, Start Date (+ Birthday when known) on the next empty row. Everything else is automatic."),
        ("Leaver", "Type their Last Working Day. AL is pro-rated and they drop off next year's summary automatically."),
        ("Birthday Leave", "Fully automatic: once a birthday is entered, 1 day of BL is applied on the birthday (or the next working day if it falls on a weekend / public holiday). To move it, type a date in 'BL date override'. No need to log BL."),
    ]),
    ("WHAT YOU ONLY CHECK", [
        ("Summary 2026 / 2027", "Balances per person. 2027 carry-forward = MIN(6, unused 2026 AL), pulled in automatically."),
        ("Calendar", "Pick Year and Month at the top. AL / MC / BL, ½ = half day, PH = public holiday, · = not employed."),
        ("⚠ Check column", "Flags mistakes in the Leave Log: outside employment dates, balance exceeded, half day across several dates, overlapping entries, etc."),
    ]),
    ("RULES BUILT IN", [
        ("Annual Leave", "12 days/year. Pro-rated for joiners & leavers: 12 × completed months ÷ 12, rounded per EA s.60E (≥ ½ day rounds up). Never below statutory 8/12/16 days (<2 / 2–5 / 5+ yrs service)."),
        ("Medical Leave", "14 days/year, not pro-rated. Auto-raised to the statutory minimum 18 days (2–5 yrs service) / 22 days (5+ yrs), EA s.60F."),
        ("Birthday Leave", "1 day, only if the birthday falls within the employment period that year."),
        ("Carry-forward", "Up to 6 unused AL days roll into the next year and must be used by 31 Dec of that year."),
        ("Public holidays", "Holidays 2026 and Holidays 2027 tabs (Selangor + national). If a moon-sighting date changes, edit it there — everything recalculates."),
    ]),
    ("COLOUR KEY", [
        ("Yellow cells", "You type here."),
        ("Grey cells", "Auto-calculated — please don't type over them."),
    ]),
]
r = 4
for head, rows in sections:
    for c in (1, 2):
        ws.cell(r, c).fill = fill(BLACK)
    ws.cell(r, 1, head).font = F(10, True, YELLOW)
    r += 1
    for k, v in rows:
        a = ws.cell(r, 1, k)
        a.font = F(10, True, BLACK)
        a.alignment = Alignment(vertical="top")
        b = ws.cell(r, 2, v)
        b.font = F(10)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        if k.startswith("Example"):
            b.fill = fill(YELLOW_SOFT)
            b.font = F(10, color="000000", italic=True)
        if k == "Yellow cells":
            a.fill = fill(YELLOW_SOFT)
        if k == "Grey cells":
            a.fill = fill(AUTO_FILL)
        r += 1
    r += 1
widths(ws, {"A": 22, "B": 120})
ws.sheet_view.showGridLines = False

# ================================================================ Settings
st = wb.create_sheet("Settings")
title(st, "Settings", "Leave policy and statutory minimums. Edit the yellow cells only.")
st["A4"] = "Today override (leave blank = real date)"
st["A5"], st["B5"] = "Effective today", '=IF(B4="",TODAY(),B4)'
st["A7"] = "COMPANY LEAVE POLICY (days per year)"
policy = [
    ("Annual Leave (AL)", 12, "Company policy (from Alfred)"),
    ("Medical Leave (MC)", 14, "Company policy (from Alfred); auto-raised to statutory tier if higher"),
    ("Birthday Leave (BL)", 1, "Company policy (from Alfred); applied automatically on the birthday"),
    ("Max AL carry-forward", 6, "Unused AL carried to next year, expires 31 Dec of that year"),
    ("New Joiner status period (months)", 3, "Status shows 'New Joiner' for this many months after start"),
    ("Low balance warning (days)", 2, "Summary turns amber at or below this AL balance"),
]
for i, (k, v, note) in enumerate(policy):
    rr = 8 + i
    st.cell(rr, 1, k)
    st.cell(rr, 2, v).fill = fill(YELLOW_SOFT)
    st.cell(rr, 3, note).font = F(9, italic=True, color=GREY_TXT)
P_AL, P_MC, P_BL, P_CF, P_NJ, P_LOW = (name(f"Policy_{k}", f"Settings!$B${8+i}")
                                       for i, k in enumerate(["AL", "MC", "BL", "MaxCF", "NewJoinerMonths", "LowBal"]))
TODAY = name("AsOfDate", "Settings!$B$5")
st["A15"] = "STATUTORY MINIMUM — Employment Act 1955 (as amended 2022), s.60E & s.60F"
header(st, 16, ["Service from (years)", "Min. AL days", "Min. MC days"], height=30)
for i, row in enumerate([(0, 8, 14), (2, 12, 18), (5, 16, 22)]):
    for j, v in enumerate(row):
        c = st.cell(17 + i, 1 + j, v)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
T_YRS, T_AL, T_MC = (name("Tier_Yrs", "Settings!$A$17:$A$19"), name("Tier_AL", "Settings!$B$17:$B$19"),
                     name("Tier_MC", "Settings!$C$17:$C$19"))
st["A21"] = ("Pro-rating: AL × completed months of service in the year ÷ 12; fraction < ½ day dropped, "
             "≥ ½ day rounded up (EA s.60E(1A)). MC is not pro-rated. Service tier is measured at 31 Dec "
             "(or last working day) of each year — the safer reading.")
st["A21"].alignment = Alignment(wrap_text=True, vertical="top")
st.merge_cells("A21:C21")
st.row_dimensions[21].height = 54
st["A23"] = "JOBSCOPER BRAND PALETTE"
for i, (k, v, fc) in enumerate([("Black", BLACK, "FFFFFF"), ("JobScoper yellow", YELLOW, BLACK),
                                ("AL", LEAVE["AL"][0], BLACK), ("MC", LEAVE["MC"][0], LEAVE["MC"][1]),
                                ("BL", LEAVE["BL"][0], LEAVE["BL"][1]), ("Public holiday", LEAVE["PH"][0], YELLOW)]):
    st.cell(24 + i, 1, k)
    c = st.cell(24 + i, 2, "#" + v)
    c.fill, c.font = fill(v), F(10, True, fc)
for rr in (7, 15, 23):
    st.cell(rr, 1).font = F(10, True, BLACK)
    st.cell(rr, 1).fill = fill(YELLOW)
st["B4"].fill = fill(YELLOW_SOFT)
st["B4"].number_format = st["B5"].number_format = "d mmm yyyy"
st["B5"].fill = fill(AUTO_FILL)
widths(st, {"A": 40, "B": 16, "C": 64})

# ================================================================ Staff
sf = wb.create_sheet("Staff")
title(sf, "Staff", "Add / remove people here only. Leavers: type Last Working Day (keep the row so history stays). Max 30 rows.")
header(sf, 4, ["Name", "Position", "Start Date", "Last Working Day\n(blank if active)", "Birthday",
               "BL date override\n(optional)", "Carry-Fwd into 2026\n(from 2025)", "Status\n(auto)", "Service\n(auto)", "Notes"])
seed = [
    ("Alfred", "People Solution Lead", date(2025, 6, 1), 1.5, "Carry-forward 1.5 days from 2025 (from old tracker)"),
    ("Shamirah", "Talent Acquisition", date(2026, 8, 24), 0, ""),
    ("Zack", "Talent Acquisition", date(2026, 11, 16), 0, "Joining 16 Nov 2026"),
]
for r in range(STAFF_FIRST, STAFF_LAST + 1):
    i = r - STAFF_FIRST
    if i < len(seed):
        n, p, s, cfw, note = seed[i]
        sf[f"A{r}"], sf[f"B{r}"], sf[f"C{r}"], sf[f"G{r}"], sf[f"J{r}"] = n, p, s, cfw, note
    lwd_past = f"AND(D{r}<>\"\",D{r}<{TODAY})"
    sf[f"H{r}"] = (f'=IF(A{r}="","",IF(C{r}="","Start date?",IF(C{r}>{TODAY},"Upcoming",'
                   f'IF({lwd_past},"Left",IF(D{r}<>"","Resigning",'
                   f'IF({TODAY}<EDATE(C{r},{P_NJ}),"New Joiner","Active"))))))')
    upto = f"IF({lwd_past},D{r},{TODAY})"
    sf[f"I{r}"] = (f'=IF(OR(A{r}="",C{r}="",C{r}>{TODAY}),"",DATEDIF(C{r},{upto},"Y")&"y "&DATEDIF(C{r},{upto},"YM")&"m")')
    for c in range(1, 11):
        cell = sf.cell(r, c)
        cell.border = BOX
        cell.font = F(10)
        cell.fill = fill(AUTO_FILL if c in (8, 9) else YELLOW_SOFT)
    for c in "CDEF":
        sf[f"{c}{r}"].number_format = "d mmm yyyy"
    sf[f"G{r}"].number_format = "0.0"
    for c in "HI":
        sf[f"{c}{r}"].alignment = Alignment(horizontal="center")
sf["E4"].comment = Comment("Date of birth (any year is fine — only day & month are used). Birthday Leave is applied automatically once entered.", "Tracker")
sf["F4"].comment = Comment("Optional. Type a date only if Birthday Leave should be taken on a different day that year.", "Tracker")
sf["G4"].comment = Comment("Only for the first year (2026). From 2027 onwards carry-forward is calculated automatically.", "Tracker")
status_cf(sf, f"H{STAFF_FIRST}:H{STAFF_LAST}", f"H{STAFF_FIRST}")
cf(sf, f"E{STAFF_FIRST}:E{STAFF_LAST}", f'AND($A{STAFF_FIRST}<>"",$E{STAFF_FIRST}="")', "FFE8B3")
dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(1940,1,1)", allow_blank=True,
                    showErrorMessage=True, error="Please enter a valid date.")
sf.add_data_validation(dv)
dv.add(f"C{STAFF_FIRST}:F{STAFF_LAST}")
dvn = DataValidation(type="custom", formula1=f"COUNTIF($A${STAFF_FIRST}:$A${STAFF_LAST},A{STAFF_FIRST})<=1",
                     allow_blank=True, showErrorMessage=True,
                     error="This name already exists. Use a unique name (e.g. add surname initial).")
sf.add_data_validation(dvn)
dvn.add(f"A{STAFF_FIRST}:A{STAFF_LAST}")
widths(sf, {"A": 18, "B": 24, "C": 14, "D": 16, "E": 14, "F": 16, "G": 15, "H": 13, "I": 10, "J": 46})
sf.freeze_panes = "B5"

# ================================================================ Holidays 2026 / 2027
HOLIDAYS = {
    2026: [
        (date(2026, 1, 1), "New Year's Day", "Selangor", "Confirmed"),
        (date(2026, 2, 1), "Thaipusam", "Selangor", "Confirmed · falls on Sunday"),
        (date(2026, 2, 2), "Thaipusam (in lieu)", "Replacement", "Confirmed"),
        (date(2026, 2, 17), "Chinese New Year", "National", "Confirmed"),
        (date(2026, 2, 18), "Chinese New Year (Day 2)", "National", "Confirmed"),
        (date(2026, 3, 7), "Nuzul Al-Quran", "Selangor", "Confirmed · falls on Saturday"),
        (date(2026, 3, 21), "Hari Raya Aidilfitri", "National", "Confirmed · declared by Keeper of the Rulers' Seal, 19 Mar 2026"),
        (date(2026, 3, 22), "Hari Raya Aidilfitri (Day 2)", "National", "Confirmed · falls on Sunday"),
        (date(2026, 3, 23), "Hari Raya Aidilfitri (in lieu)", "Replacement", "Confirmed"),
        (date(2026, 5, 1), "Labour Day", "National", "Confirmed"),
        (date(2026, 5, 27), "Hari Raya Haji (Aidiladha)", "National", "Confirmed · declared by Keeper of the Rulers' Seal"),
        (date(2026, 5, 31), "Wesak Day", "National", "Confirmed · falls on Sunday"),
        (date(2026, 6, 1), "Agong's Birthday", "National", "Confirmed"),
        (date(2026, 6, 2), "Wesak Day (in lieu)", "Replacement", "Confirmed · 1 Jun already a holiday"),
        (date(2026, 6, 17), "Awal Muharram", "National", "Confirmed"),
        (date(2026, 8, 25), "Maulidur Rasul", "National", "Confirmed · malaysia.gov.my calendar"),
        (date(2026, 8, 31), "National Day", "National", "Confirmed"),
        (date(2026, 9, 16), "Malaysia Day", "National", "Confirmed"),
        (date(2026, 10, 25), "Sultan of Selangor Silver Jubilee", "Selangor", "Confirmed · special holiday, falls on Sunday"),
        (date(2026, 10, 26), "Silver Jubilee (in lieu)", "Replacement", "Confirmed · Selangor State Secretary, Sep 2026"),
        (date(2026, 11, 8), "Deepavali", "National", "Confirmed · falls on Sunday"),
        (date(2026, 11, 9), "Deepavali (in lieu)", "Replacement", "Confirmed"),
        (date(2026, 12, 11), "Sultan of Selangor's Birthday", "Selangor", "Confirmed"),
        (date(2026, 12, 25), "Christmas Day", "National", "Confirmed"),
    ],
    2027: [
        (date(2027, 1, 1), "New Year's Day", "Selangor", "Fixed date"),
        (date(2027, 1, 22), "Thaipusam", "Selangor", "Scheduled"),
        (date(2027, 2, 6), "Chinese New Year", "National", "Scheduled · falls on Saturday"),
        (date(2027, 2, 7), "Chinese New Year (Day 2)", "National", "Scheduled · falls on Sunday"),
        (date(2027, 2, 8), "Chinese New Year (in lieu)", "Replacement", "Scheduled"),
        (date(2027, 2, 24), "Nuzul Al-Quran", "Selangor", "Subject to moon sighting"),
        (date(2027, 3, 10), "Hari Raya Aidilfitri", "National", "Subject to moon sighting"),
        (date(2027, 3, 11), "Hari Raya Aidilfitri (Day 2)", "National", "Subject to moon sighting"),
        (date(2027, 5, 1), "Labour Day", "National", "Fixed date · falls on Saturday"),
        (date(2027, 5, 17), "Hari Raya Haji (Aidiladha)", "National", "Subject to moon sighting"),
        (date(2027, 5, 20), "Wesak Day", "National", "Scheduled"),
        (date(2027, 6, 6), "Awal Muharram", "National", "Subject to moon sighting · falls on Sunday"),
        (date(2027, 6, 7), "Agong's Birthday", "National", "Fixed (first Monday of June)"),
        (date(2027, 6, 8), "Awal Muharram (in lieu)", "Replacement", "Subject to moon sighting · 7 Jun already a holiday"),
        (date(2027, 8, 15), "Maulidur Rasul", "National", "Subject to moon sighting · falls on Sunday"),
        (date(2027, 8, 16), "Maulidur Rasul (in lieu)", "Replacement", "Subject to moon sighting"),
        (date(2027, 8, 31), "National Day", "National", "Fixed date"),
        (date(2027, 9, 16), "Malaysia Day", "National", "Fixed date"),
        (date(2027, 10, 28), "Deepavali", "National", "Scheduled"),
        (date(2027, 12, 11), "Sultan of Selangor's Birthday", "Selangor", "Fixed date · falls on Saturday"),
        (date(2027, 12, 25), "Christmas Day", "National", "Fixed date · falls on Saturday"),
    ],
}
for y in YEARS:
    hd = wb.create_sheet(f"Holidays {y}")
    title(hd, f"Public Holidays {y} — Selangor + National",
          "Weekends are already excluded automatically. If a moon-sighting date changes, edit the date here — everything recalculates.")
    header(hd, 4, ["Date", "Holiday", "Type", "Status / source"])
    for r in range(HOL_FIRST, HOL_LAST + 1):
        i = r - HOL_FIRST
        if i < len(HOLIDAYS[y]):
            for j, v in enumerate(HOLIDAYS[y][i]):
                hd.cell(r, 1 + j, v)
        for c in range(1, 5):
            cell = hd.cell(r, c)
            cell.fill = fill(YELLOW_SOFT)
            cell.border = BOX
            cell.font = F(10)
        hd.cell(r, 1).number_format = "ddd, d mmm yyyy"
    cf(hd, f"D{HOL_FIRST}:D{HOL_LAST}", f'ISNUMBER(SEARCH("moon",$D{HOL_FIRST}))', None, "B54708")
    cf(hd, f"A{HOL_FIRST}:B{HOL_LAST}", f'AND($A{HOL_FIRST}<>"",WEEKDAY($A{HOL_FIRST},2)>5)', None, "8A8A8A")
    dvh = DataValidation(type="date", operator="between", formula1=f"DATE({y},1,1)", formula2=f"DATE({y},12,31)",
                         allow_blank=True, showErrorMessage=True, error=f"Date must be in {y}.")
    hd.add_data_validation(dvh)
    dvh.add(f"A{HOL_FIRST}:A{HOL_LAST}")
    widths(hd, {"A": 20, "B": 36, "C": 14, "D": 60})
    hd.freeze_panes = "A5"

# ================================================================ _calc (hidden helper, one row per Staff row)
cs = wb.create_sheet("_calc")
cs["A1"] = "Helper calculations — do not edit."
cs["A4"] = "Name"
for r in range(STAFF_FIRST, STAFF_LAST + 1):
    cs[f"A{r}"] = f'=IF(Staff!A{r}="","",Staff!A{r})'

BLOCK = ["InYr", "Rank", "EmpStart", "EmpEnd", "Months", "SvcYrs", "ALbase", "ALent", "CF", "ALtaken", "ALrem",
         "MCent", "MCtaken", "MCrem", "Bday", "BLent", "AutoBL", "BLtaken", "BLrem", "Note"]
CALC = {}


def build_year_block(start_col, year, prev=None):
    cols = {k: L(start_col + i) for i, k in enumerate(BLOCK)}
    cs.cell(1, start_col, f"Year block {year}")
    cs.cell(2, start_col, year)
    Y = f"${L(start_col)}$2"
    for i, k in enumerate(BLOCK):
        cs.cell(4, start_col + i, k)
    for r in range(STAFF_FIRST, STAFF_LAST + 1):
        c = {k: f"{v}{r}" for k, v in cols.items()}
        s = lambda col: f"Staff!${col}{r}"  # noqa: E731
        f = {}
        f["InYr"] = (f'=AND($A{r}<>"",{s(S_START)}<>"",{s(S_START)}<=DATE({Y},12,31),'
                     f'OR({s(S_LWD)}="",{s(S_LWD)}>=DATE({Y},1,1)))')
        f["Rank"] = f'=IF({c["InYr"]},COUNTIF({cols["InYr"]}${STAFF_FIRST}:{c["InYr"]},TRUE),"")'
        f["EmpStart"] = f'=IF({c["InYr"]},MAX({s(S_START)},DATE({Y},1,1)),"")'
        f["EmpEnd"] = f'=IF({c["InYr"]},IF({s(S_LWD)}="",DATE({Y},12,31),MIN({s(S_LWD)},DATE({Y},12,31))),"")'
        f["Months"] = f'=IF({c["InYr"]},DATEDIF({c["EmpStart"]},{c["EmpEnd"]}+1,"M"),0)'
        f["SvcYrs"] = f'=IF({c["InYr"]},DATEDIF({s(S_START)},{c["EmpEnd"]},"Y"),0)'
        f["ALbase"] = f'=IF({c["InYr"]},MAX({P_AL},INDEX({T_AL},MATCH({c["SvcYrs"]},{T_YRS},1))),0)'
        f["ALent"] = f'=IF({c["InYr"]},ROUND({c["ALbase"]}*{c["Months"]}/12,0),0)'
        if prev is None:
            f["CF"] = f'=IF({c["InYr"]},N({s(S_CF)}),0)'
        else:
            f["CF"] = f'=IF(AND({c["InYr"]},{prev["InYr"]}{r}),MIN({P_CF},MAX(0,{prev["ALrem"]}{r})),0)'
        for t in ("AL", "MC"):
            f[f"{t}taken"] = (f'=IF($A{r}="",0,SUMIFS({log("H")},{log("A")},$A{r},{log("B")},"{t}",'
                              f'{log("I")},{Y},{log("F")},"<>Cancelled"))')
        f["ALrem"] = f'={c["ALent"]}+{c["CF"]}-{c["ALtaken"]}'
        f["MCent"] = f'=IF({c["InYr"]},MAX({P_MC},INDEX({T_MC},MATCH({c["SvcYrs"]},{T_YRS},1))),0)'
        f["MCrem"] = f'={c["MCent"]}-{c["MCtaken"]}'
        # Birthday Leave: automatic on the birthday, or next working day if weekend/PH; override wins
        f["Bday"] = f'=IF(AND({c["InYr"]},{s(S_BDAY)}<>""),DATE({Y},MONTH({s(S_BDAY)}),DAY({s(S_BDAY)})),"")'
        f["BLent"] = (f'=IF({c["Bday"]}="",0,IF(AND({c["Bday"]}>={c["EmpStart"]},{c["Bday"]}<={c["EmpEnd"]}),{P_BL},0))')
        f["AutoBL"] = (f'=IF({c["BLent"]}=0,"",IF(AND({s(S_BLOVR)}<>"",YEAR({s(S_BLOVR)})={Y}),{s(S_BLOVR)},'
                       f'WORKDAY({c["Bday"]}-1,1,{HOLS[year]})))')
        f["BLtaken"] = f'=IF({c["AutoBL"]}="",0,IF({c["AutoBL"]}<={TODAY},{c["BLent"]},0))'
        f["BLrem"] = f'={c["BLent"]}-{c["BLtaken"]}'
        bl_note = (f'IF({s(S_BDAY)}="","Add birthday in Staff for BL. ",IF({c["BLent"]}=0,"Birthday outside employment period, no BL. ",'
                   f'IF({c["BLtaken"]}>0,"BL taken "&TEXT({c["AutoBL"]},"d mmm")&". ","BL auto-scheduled "&TEXT({c["AutoBL"]},"ddd d mmm")&". ")))')
        f["Note"] = (f'=IF(NOT({c["InYr"]}),"",TRIM('
                     f'IF({c["Months"]}<12,"AL pro-rated: "&{c["Months"]}&"/12 completed months. ","")'
                     f'&IF({c["ALbase"]}>{P_AL},"AL raised to statutory "&{c["ALbase"]}&" days. ","")'
                     f'&IF({c["MCent"]}>{P_MC},"MC raised to statutory "&{c["MCent"]}&" days (EA s.60F). ","")'
                     f'&IF({c["CF"]}>0,"Carry-fwd "&{c["CF"]}&" day(s), use by 31 Dec {year}. ","")'
                     f'&{bl_note}'
                     f'&IF({s(S_LWD)}<>"","Last working day "&TEXT({s(S_LWD)},"d mmm yyyy")&". ","")))')
        for k in BLOCK:
            cs[c[k]] = f[k]
    rng = {k: name(f"{k}_{year}", f"_calc!${v}${STAFF_FIRST}:${v}${STAFF_LAST}") for k, v in cols.items()}
    rng["_cols"] = cols
    return rng


CALC[2026] = build_year_block(2, 2026)
CALC[2027] = build_year_block(2 + len(BLOCK), 2027, prev=CALC[2026]["_cols"])
for k in ("Bday", "AutoBL", "EmpStart", "EmpEnd"):
    for y in YEARS:
        col = CALC[y]["_cols"][k]
        for r in range(STAFF_FIRST, STAFF_LAST + 1):
            cs[f"{col}{r}"].number_format = "d mmm yyyy"

# calendar month block
cm = 2 + 2 * len(BLOCK) + 1
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
cs.cell(1, cm, "Calendar block")
cs.cell(4, cm, "InMonth")
cs.cell(4, cm + 1, "RankMonth")
MCOL = L(cm + 3)
cs[f"{MCOL}4"] = "Months"
for i, m in enumerate(MONTHS):
    cs[f"{MCOL}{5+i}"] = m
cs.cell(2, cm, "=Calendar!$B$2")
cs.cell(2, cm + 1, f"=MATCH(Calendar!$G$2,${MCOL}$5:${MCOL}$16,0)")
CY, CM = name("CalYear", f"_calc!${L(cm)}$2"), name("CalMonth", f"_calc!${L(cm+1)}$2")
ci, cr = L(cm), L(cm + 1)
for r in range(STAFF_FIRST, STAFF_LAST + 1):
    cs[f"{ci}{r}"] = (f'=AND($A{r}<>"",Staff!${S_START}{r}<>"",Staff!${S_START}{r}<=EOMONTH(DATE({CY},{CM},1),0),'
                      f'OR(Staff!${S_LWD}{r}="",Staff!${S_LWD}{r}>=DATE({CY},{CM},1)))')
    cs[f"{cr}{r}"] = f'=IF({ci}{r},COUNTIF(${ci}${STAFF_FIRST}:{ci}{r},TRUE),"")'
CAL_RANK = name("CalRank", f"_calc!${cr}${STAFF_FIRST}:${cr}${STAFF_LAST}")
KC, KT = L(cm + 5), L(cm + 6)
cs[f"{KC}4"], cs[f"{KT}4"] = "Code", "Label"
codes = [(0.5, "AL½"), (1, "AL"), (5, "MC½"), (10, "MC"), (5.5, "AL½+MC½")]
for i, (k, v) in enumerate(codes):
    cs[f"{KC}{5+i}"], cs[f"{KT}{5+i}"] = k, v
CODE_K = name("Code_Keys", f"_calc!${KC}$5:${KC}${4+len(codes)}")
CODE_T = name("Code_Labels", f"_calc!${KT}$5:${KT}${4+len(codes)}")
cs.sheet_state = "hidden"

# ================================================================ Leave Log
lg = wb.create_sheet("Leave Log")
title(lg, "Leave Log", "Apply leave here: one row per leave. Type in yellow columns only; grey columns calculate automatically. "
                       "Leave 'To' blank for a single day. Birthday Leave is automatic — no need to log it.")
header(lg, 4, ["Name", "Type", "From", "To", "Day", "Status", "Remarks", "Working\nDays (auto)", "Year\n(auto)",
               "⚠ Check (auto)", "_idx", "_remaining", "_to", "_code", "_bl", "_issues"])
migrated = [
    ("Alfred", "AL", date(2026, 1, 2), None, "Full", "Approved", "Migrated from old tracker"),
    ("Alfred", "AL", date(2026, 4, 29), date(2026, 4, 30), "Full", "Approved", "Migrated from old tracker"),
    ("Alfred", "AL", date(2026, 6, 30), None, "Full", "Approved", "Migrated from old tracker"),
]
for r in range(LOG_FIRST, LOG_LAST + 1):
    i = r - LOG_FIRST
    if i < len(migrated):
        for j, v in enumerate(migrated[i]):
            lg.cell(r, 1 + j, v)
    half = f'OR(E{r}="AM",E{r}="PM")'
    hy = hols_for(f"C{r}")
    bl_in = f'AND(ISNUMBER(O{r}),O{r}>=C{r},O{r}<=M{r})'
    # H working days (the person's auto Birthday Leave day is never deducted as AL/MC)
    lg[f"H{r}"] = (f'=IF(OR(A{r}="",C{r}=""),"",IF({half},IF({bl_in},0,0.5*NETWORKDAYS(C{r},C{r},{hy})),'
                   f'NETWORKDAYS(C{r},M{r},{hy})-IF({bl_in},NETWORKDAYS(O{r},O{r},{hy}),0)))')
    lg[f"I{r}"] = f'=IF(C{r}="","",YEAR(C{r}))'
    lg[f"K{r}"] = f'=IF(A{r}="","",IFERROR(MATCH(A{r},{staff("A")},0),""))'

    def rem(y):
        b = CALC[y]
        return f'IF(B{r}="AL",INDEX({b["ALrem"]},K{r}),IF(B{r}="MC",INDEX({b["MCrem"]},K{r}),""))'
    lg[f"L{r}"] = f'=IF(OR(K{r}="",I{r}=""),"",IFERROR(IF(I{r}=2026,{rem(2026)},IF(I{r}=2027,{rem(2027)},"")),""))'
    lg[f"M{r}"] = f'=IF(C{r}="","",IF(D{r}="",C{r},D{r}))'
    lg[f"N{r}"] = (f'=IF(OR(A{r}="",C{r}=""),"",IF(B{r}="AL",1,IF(B{r}="MC",10,0))*IF({half},0.5,1))')
    lg[f"O{r}"] = (f'=IF(OR(K{r}="",C{r}=""),"",IFERROR(IF(YEAR(C{r})=2026,INDEX({CALC[2026]["AutoBL"]},K{r}),'
                   f'IF(YEAR(C{r})=2027,INDEX({CALC[2027]["AutoBL"]},K{r}),"")),""))')
    st_ = f"INDEX({staff(S_START)},K{r})"
    lwd = f"INDEX({staff(S_LWD)},K{r})"
    checks = [
        f'IF(B{r}="","Pick a leave type. ","")',
        f'IF(C{r}="","Enter From date. ","")',
        f'IF(AND(D{r}<>"",C{r}<>"",D{r}<C{r}),"To is before From. ","")',
        f'IF(AND(C{r}<>"",C{r}<{st_}),"Before start date. ","")',
        f'IF(AND(C{r}<>"",{lwd}<>""),IF(M{r}>{lwd},"After last working day. ",""),"")',
        f'IF(AND(C{r}<>"",H{r}=0),"No working days in range (weekend/PH/BL). ","")',
        f'IF(AND({half},D{r}<>"",D{r}<>C{r}),"Half day must be a single date. ","")',
        f'IF(AND(C{r}<>"",YEAR(M{r})<>YEAR(C{r})),"Crosses year end, split into 2 rows. ","")',
        f'IF(AND(I{r}<>"",I{r}<>2026,I{r}<>2027),"No summary for this year yet. ","")',
        f'IF(AND(L{r}<>"",N(L{r})<0),"Balance exceeded. ","")',
        (f'IF(AND(C{r}<>"",COUNTIFS({log("A")},A{r},{log("C")},"<="&M{r},{log("M")},">="&C{r},'
         f'{log("F")},"<>Cancelled")>1),"Overlaps another entry. ","")'),
        f'IF(F{r}="Pending","Pending approval. ","")',
    ]
    lg[f"P{r}"] = (f'=IF(OR(A{r}="",F{r}="Cancelled"),"",IF(K{r}="","Name not in Staff list.",'
                   f'IFERROR(TRIM({"&".join(checks)}),"Check dates.")))')
    lg[f"J{r}"] = (f'=IF(A{r}="","",IF(F{r}="Cancelled","Cancelled",IF(P{r}="","✓ OK"&IF(AND(C{r}<>"",{bl_in}),'
                   f'" · Birthday Leave day not deducted",""),P{r})))')
    for c in range(1, 17):
        cell = lg.cell(r, c)
        cell.border = BOX
        cell.font = F(10)
        cell.fill = fill(YELLOW_SOFT if c <= 7 else AUTO_FILL)
    for c in "CDMO":
        lg[f"{c}{r}"].number_format = "ddd, d mmm yyyy"
    for c in "BEFHI":
        lg[f"{c}{r}"].alignment = Alignment(horizontal="center")
widths(lg, {"A": 16, "B": 7, "C": 17, "D": 17, "E": 7, "F": 11, "G": 30, "H": 10, "I": 7, "J": 60,
            "K": 6, "L": 8, "M": 14, "N": 6, "O": 14, "P": 30})
for c in "KLMNOP":
    lg.column_dimensions[c].hidden = True
lg.freeze_panes = "B5"
for rng, f1 in ((f"A{LOG_FIRST}:A{LOG_LAST}", f"=Staff!$A${STAFF_FIRST}:$A${STAFF_LAST}"),
                (f"B{LOG_FIRST}:B{LOG_LAST}", '"AL,MC"'), (f"E{LOG_FIRST}:E{LOG_LAST}", '"Full,AM,PM"'),
                (f"F{LOG_FIRST}:F{LOG_LAST}", '"Approved,Pending,Cancelled"')):
    d = DataValidation(type="list", formula1=f1, allow_blank=True, showErrorMessage=True,
                       error="Please pick from the dropdown list.")
    lg.add_data_validation(d)
    d.add(rng)
dvd = DataValidation(type="date", operator="greaterThan", formula1="DATE(2025,1,1)", allow_blank=True,
                     showErrorMessage=True, error="Please enter a valid date.")
lg.add_data_validation(dvd)
dvd.add(f"C{LOG_FIRST}:D{LOG_LAST}")
for t in ("AL", "MC"):
    cf(lg, f"B{LOG_FIRST}:B{LOG_LAST}", f'$B{LOG_FIRST}="{t}"', LEAVE[t][0], LEAVE[t][1], True)
j0 = f"$J{LOG_FIRST}"
cf(lg, f"J{LOG_FIRST}:J{LOG_LAST}", f'LEFT({j0},4)="✓ OK"', None, "1E6B34")
cf(lg, f"J{LOG_FIRST}:J{LOG_LAST}", f'LEFT({j0},7)="Pending"', None, "8A5A00")
cf(lg, f"J{LOG_FIRST}:J{LOG_LAST}", f'AND({j0}<>"",LEFT({j0},4)<>"✓ OK",{j0}<>"Cancelled")', "FDECEA", "B42318", True)
cf(lg, f"A{LOG_FIRST}:J{LOG_LAST}", f'$F{LOG_FIRST}="Cancelled"', None, "ADB5BD")

# ================================================================ Summaries
SUM_FIRST, SUM_LAST = 6, 35
SUM_COLS = [("#", 4), ("Name", 16), ("Position", 22), ("Status", 12), ("Start Date", 12), ("Last Working Day", 12),
            ("AL Entitled\n(pro-rated)", 11), ("Carry-Fwd", 9), ("AL Total", 9), ("AL Taken", 9), ("AL Balance", 10),
            ("MC Entitled", 10), ("MC Taken", 9), ("MC Balance", 10), ("BL Entitled", 9), ("BL Date\n(auto)", 12),
            ("BL Taken", 9), ("AL Used %", 9), ("Notes (auto)", 62), ("_idx", 5)]


def build_summary(year):
    s = wb.create_sheet(f"Summary {year}")
    b = CALC[year]
    title(s, f"Leave Summary {year}", "Auto-calculated from Staff + Leave Log — nothing to type here. One row per person employed during the year.")
    kpis = [("B3", "Headcount", f'=COUNTIF({b["InYr"]},TRUE)'),
            ("E3", "On leave", (f'=IF(YEAR({TODAY})<>{year},"–",COUNTIFS({log("C")},"<="&{TODAY},{log("M")},">="&{TODAY},'
                                      f'{log("F")},"<>Cancelled")+COUNTIF({b["AutoBL"]},{TODAY}))')),
            ("H3", "Joiners", f'=COUNTIF({staff(S_STATUS)},"Upcoming")'),
            ("K3", "As of", f"={TODAY}")]
    for ref, lab, f in kpis:
        c = s[ref]
        c.value = lab
        c.font = F(9, True, BLACK)
        c.fill = fill(YELLOW)
        c.alignment = Alignment(horizontal="center", vertical="center")
        v = c.offset(0, 1)
        v.value = f
        v.font = F(12, True, YELLOW)
        v.fill = fill(BLACK)
        v.alignment = Alignment(horizontal="center", vertical="center")
    s["L3"].number_format = "d mmm yyyy"
    s.merge_cells("L3:M3")
    s.row_dimensions[3].height = 26
    header(s, 5, [c for c, _ in SUM_COLS], height=36)
    for i, (_, w) in enumerate(SUM_COLS):
        s.column_dimensions[L(i + 1)].width = w
    s.column_dimensions["T"].hidden = True
    for r in range(SUM_FIRST, SUM_LAST + 1):
        s.cell(r, 1, r - SUM_FIRST + 1)
        s[f"T{r}"] = f'=IFERROR(MATCH($A{r},{b["Rank"]},0),"")'

        def ix(rng):
            return f'=IF($T{r}="","",INDEX({rng},$T{r}))'
        vals = [ix(staff(S_NAME)), ix(staff(S_POS)), ix(staff(S_STATUS)), ix(staff(S_START)),
                f'=IF($T{r}="","",IF(INDEX({staff(S_LWD)},$T{r})="","",INDEX({staff(S_LWD)},$T{r})))',
                ix(b["ALent"]), ix(b["CF"]), f'=IF($T{r}="","",G{r}+H{r})', ix(b["ALtaken"]), ix(b["ALrem"]),
                ix(b["MCent"]), ix(b["MCtaken"]), ix(b["MCrem"]), ix(b["BLent"]), ix(b["AutoBL"]), ix(b["BLtaken"]),
                f'=IF(OR($T{r}="",N(I{r})=0),"",J{r}/I{r})', ix(b["Note"])]
        for k, v in enumerate(vals):
            s.cell(r, 2 + k, v)
        for c in range(1, 20):
            cell = s.cell(r, c)
            cell.font = F(10, color="ADB5BD" if c == 1 else "212529")
            cell.alignment = Alignment(horizontal="left" if c in (2, 3, 19) else "center", vertical="center",
                                       wrap_text=(c == 19))
            cell.border = Border(bottom=Side(style="thin", color="EEEEEE"))
        for c in "EFP":
            s[f"{c}{r}"].number_format = "d mmm yyyy"
        s[f"R{r}"].number_format = "0%"
        s[f"B{r}"].font = F(10, True, BLACK)
    s["T3"] = f"={P_LOW}"   # local copy: Sheets conditional formatting can't reference other tabs
    last = SUM_LAST + 2
    s.cell(last, 2, "Team total").font = F(10, True)
    for c in "JMQ":
        s[f"{c}{last}"] = f"=SUM({c}{SUM_FIRST}:{c}{SUM_LAST})"
        s[f"{c}{last}"].font = F(10, True)
        s[f"{c}{last}"].alignment = Alignment(horizontal="center")
    status_cf(s, f"D{SUM_FIRST}:D{SUM_LAST}", f"D{SUM_FIRST}")
    for colL in ("K", "N"):
        a = f"{colL}{SUM_FIRST}"
        cf(s, f"{colL}{SUM_FIRST}:{colL}{SUM_LAST}", f'AND($B{SUM_FIRST}<>"",N({a})<0)', "F8D7DA", "8A1C24", True)
    a = f"K{SUM_FIRST}"
    cf(s, f"K{SUM_FIRST}:K{SUM_LAST}", f'AND($B{SUM_FIRST}<>"",N({a})>=0,N({a})<=$T$3)', "FFE8B3", "8A5A00", True)
    cf(s, f"K{SUM_FIRST}:K{SUM_LAST}", f'AND($B{SUM_FIRST}<>"",N({a})>$T$3)', "DFF3E4", "1E6B34", True)
    cf(s, f"R{SUM_FIRST}:R{SUM_LAST}", f'AND($R{SUM_FIRST}<>"",$R{SUM_FIRST}>=0.8)', None, "B42318", True)
    cf(s, f"P{SUM_FIRST}:Q{SUM_LAST}", f'AND($B{SUM_FIRST}<>"",$O{SUM_FIRST}>0)', LILAC, LEAVE["BL"][1])
    cf(s, f"A{SUM_FIRST}:S{SUM_LAST}", f'AND($B{SUM_FIRST}<>"",MOD(ROW(),2)=0)', "FAFAFA")   # lowest priority
    for cols_, (fc, tc) in (("GHIJK", (LEAVE["AL"][0], BLACK)), ("LMN", (SKY, LEAVE["MC"][1])), ("OPQ", (LILAC, LEAVE["BL"][1]))):
        for colL in cols_:
            s[f"{colL}5"].fill = fill(fc)
            s[f"{colL}5"].font = F(10, True, tc)
    s.freeze_panes = "C6"
    s.sheet_view.showGridLines = False
    return s


for y in YEARS:
    build_summary(y)

# ================================================================ Calendar
cal = wb.create_sheet("Calendar")
title(cal, "Leave Calendar")
cal["A2"], cal["B2"] = "Year ▸", 2026
cal["E2"], cal["G2"] = "Month ▸", "Oct"
for ref in ("A2", "E2"):
    cal[ref].font = F(10, True, BLACK)
    cal[ref].alignment = Alignment(horizontal="right")
for ref in ("B2", "G2"):
    cal[ref].fill = fill(YELLOW)
    cal[ref].font = F(12, True, BLACK)
    cal[ref].border = BOX
    cal[ref].alignment = Alignment(horizontal="center")
cal.merge_cells("B2:C2")
cal.merge_cells("E2:F2")
cal.merge_cells("G2:I2")
for ref, f1 in (("B2", '"2026,2027"'), ("G2", '"' + ",".join(MONTHS) + '"')):
    d = DataValidation(type="list", formula1=f1, allow_blank=False)
    cal.add_data_validation(d)
    d.add(ref)
col = 12
for code, lab in [("AL", "Annual"), ("MC", "Medical"), ("BL", "Birthday (auto)"), ("PH", "Public holiday")]:
    fc, tc = LEAVE[code]
    c = cal.cell(2, col, code)
    c.fill, c.font, c.alignment, c.border = fill(fc), F(9, True, tc), Alignment(horizontal="center"), BOX
    cal.cell(2, col + 1, lab).font = F(9)
    col += 4
cal.cell(2, col, "½ = half day · · = not employed").font = F(9, italic=True, color=GREY_TXT)

CR_DATE, CR_DOW, CR_PH = 4, 5, 6
P_FIRST, P_LAST = 7, 36
DAY1, DAYN = 2, 32           # B..AF
TOT = {"AL": 33, "MC": 34, "BL": 35}
IDX, BLC = 36, 37            # hidden: staff index, this person's auto BL date
header(cal, CR_DATE, ["Name"], 1)
cal.cell(CR_PH, 1, "Public holiday").font = F(8, italic=True, color=GREY_TXT)
cal.cell(CR_PH, 1).alignment = Alignment(vertical="bottom")
for d in range(1, 32):
    c = DAY1 + d - 1
    colL = L(c)
    D = f"{colL}${CR_DATE}"
    cal.cell(CR_DATE, c, f'=IF({d}>DAY(EOMONTH(DATE($B$2,{CM},1),0)),"",DATE($B$2,{CM},{d}))').number_format = "d"
    cal.cell(CR_DOW, c, f'=IF({D}="","",{D})').number_format = "ddd"
    cal.cell(CR_PH, c, (f'=IF({D}="","",IFERROR(INDEX({HOL_NAMES[2026]},MATCH({D},{HOLS[2026]},0)),'
                        f'IFERROR(INDEX({HOL_NAMES[2027]},MATCH({D},{HOLS[2027]},0)),"")))'))
    cal.cell(CR_PH, c).alignment = Alignment(text_rotation=90, vertical="bottom", horizontal="center")
    cal.cell(CR_PH, c).font = F(7, color="000000")
    cal.column_dimensions[colL].width = 4.8
    h = cal.cell(CR_DATE, c)
    h.font, h.fill, h.alignment = F(9, True, "FFFFFF"), fill(BLACK), Alignment(horizontal="center", vertical="center")
    w = cal.cell(CR_DOW, c)
    w.font, w.alignment = F(8, color=GREY_TXT), Alignment(horizontal="center")
for k, c in TOT.items():
    h = cal.cell(CR_DATE, c, f"{k}\ndays")
    h.font = F(9, True, LEAVE[k][1])
    h.fill = fill(LEAVE[k][0])
    h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cal.column_dimensions[L(c)].width = 6.5
cal.row_dimensions[CR_PH].height = 95
cal.column_dimensions["A"].width = 16
cal.column_dimensions[L(IDX)].hidden = True
cal.column_dimensions[L(BLC)].hidden = True

for r in range(P_FIRST, P_LAST + 1):
    n = r - P_FIRST + 1
    ix = f"${L(IDX)}{r}"
    bl = f"${L(BLC)}{r}"
    cal.cell(r, IDX, f'=IFERROR(MATCH({n},{CAL_RANK},0),"")')
    cal.cell(r, BLC, (f'=IF({ix}="","",IFERROR(IF({CY}=2026,INDEX({CALC[2026]["AutoBL"]},{ix}),'
                      f'INDEX({CALC[2027]["AutoBL"]},{ix})),""))'))
    cal.cell(r, 1, f'=IF({ix}="","",INDEX({staff(S_NAME)},{ix}))').font = F(10, True, BLACK)
    cal.cell(r, 1).border = BOX
    for c in range(DAY1, DAYN + 1):
        D = f"{L(c)}${CR_DATE}"
        code = f'SUMIFS({log("N")},{log("A")},$A{r},{log("C")},"<="&{D},{log("M")},">="&{D},{log("F")},"<>Cancelled")'
        f = (f'=IF(OR($A{r}="",{D}=""),"",IF(OR({D}<INDEX({staff(S_START)},{ix}),AND(INDEX({staff(S_LWD)},{ix})<>"",'
             f'{D}>INDEX({staff(S_LWD)},{ix}))),"·",IF(WEEKDAY({D},2)>5,"",IF({is_ph(D)},"PH",IF({D}={bl},"BL",'
             f'IF({code}=0,"",IFERROR(INDEX({CODE_T},MATCH({code},{CODE_K},0)),"⚠")))))))')
        cell = cal.cell(r, c, f)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.font = F(8, True)
        cell.border = BOX
    rowr = f"$B{r}:$AF{r}"
    for k, c in TOT.items():
        if k == "BL":
            fx = f'=IF($A{r}="","",COUNTIF({rowr},"BL"))'
        else:
            fx = (f'=IF($A{r}="","",COUNTIF({rowr},"{k}")+0.5*COUNTIF({rowr},"{k}½")'
                  f'+0.5*COUNTIF({rowr},"{k}½+*")+0.5*COUNTIF({rowr},"*+{k}½"))')
        cal.cell(r, c, fx)
        cal.cell(r, c).alignment = Alignment(horizontal="center")
        cal.cell(r, c).font = F(10, True, LEAVE[k][1] if k != "AL" else BLACK)
grid = f"B{P_FIRST}:AF{P_LAST}"
g0 = f"B{P_FIRST}"
cf(cal, grid, f'$A{P_FIRST}=""', "FFFFFF")
cf(cal, grid, f'{g0}="·"', "F7F7F7", "CED4DA")
cf(cal, grid, f'{g0}="PH"', LEAVE["PH"][0], LEAVE["PH"][1])
for k in ("AL", "MC", "BL"):
    cf(cal, grid, f'LEFT({g0},2)="{k}"', LEAVE[k][0], LEAVE[k][1])
cf(cal, grid, f'{g0}="⚠"', "B42318", "FFFFFF")
cf(cal, grid, f'AND(B${CR_DATE}<>"",WEEKDAY(B${CR_DATE},2)>5)', WEEKEND)
cf(cal, f"B{CR_DATE}:AF{CR_DATE}", f'AND(B${CR_DATE}<>"",B${CR_DATE}=$AJ$2)', YELLOW, BLACK)
cf(cal, f"B{CR_DOW}:AF{CR_DOW}", f'AND(B${CR_DATE}<>"",WEEKDAY(B${CR_DATE},2)>5)', WEEKEND)
cf(cal, f"B{CR_PH}:AF{CR_PH}", f'B{CR_PH}<>""', "FFF3B0")
cal["AJ2"] = f"={TODAY}"   # local copy for conditional formatting
cal.freeze_panes = f"B{P_FIRST}"
cal.sheet_view.showGridLines = False

# ---------------------------------------------------------------- finishing
order = ["How to use", "Summary 2026", "Summary 2027", "Calendar", "Leave Log", "Staff",
         "Holidays 2026", "Holidays 2027", "Settings", "_calc"]
wb._sheets = [wb[n] for n in order]
wb.active = 1
TABS = {"How to use": YELLOW, "Summary 2026": BLACK, "Summary 2027": BLACK, "Calendar": BLACK,
        "Leave Log": YELLOW, "Staff": YELLOW, "Holidays 2026": "BFBFBF", "Holidays 2027": "BFBFBF", "Settings": "BFBFBF"}
for n, colr in TABS.items():
    wb[n].sheet_properties.tabColor = colr
for w in wb.worksheets:
    for row in w.iter_rows():
        for c in row:
            if c.font is None or c.font.name in (None, "Calibri"):
                f0 = c.font
                c.font = Font(name=FONT, size=f0.size or 10, bold=f0.bold, italic=f0.italic, color=f0.color)
    w.page_setup.orientation = "landscape"
    w.page_setup.paperSize = w.PAPERSIZE_A4
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.page_setup.fitToWidth = 1
    w.page_setup.fitToHeight = 0
for n, ref in NAMES.items():
    wb.defined_names[n] = DefinedName(n, attr_text=ref)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
